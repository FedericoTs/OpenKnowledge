"""Spreadsheets.

A spreadsheet is the format most likely to hold the exact numbers an employee
asks about - grade bands, per-country limits, approval matrices - and the one a
generic text extractor mangles worst, because a row means nothing without its
header.

Every row becomes one labelled block with a real cell locator, so a citation
points at ``Limits!A7`` rather than at the file.

Formulas are read as their last computed value, not their source: an employee
asking about the travel cap wants ``EUR 500``, not ``=B2*1.1``. That value comes
from whatever the spreadsheet last saved, which is a caveat worth knowing but
still far better than indexing the formula text.
"""

from __future__ import annotations

import contextlib
import datetime as _dt
import logging

from .blocks import (
    Block,
    BlockKind,
    ParsedDocument,
    looks_like_header_row,
    normalise,
    table_row_text,
)

log = logging.getLogger(__name__)

#: Bound a runaway sheet. Someone's 200k-row export is data, not documentation,
#: and indexing it would swamp the corpus with rows nobody will ever ask about.
MAX_ROWS_PER_SHEET = 2_000


def _render(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, _dt.datetime):
        return value.date().isoformat()
    if isinstance(value, _dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalise(str(value))


def parse_xlsx(data: bytes, *, title: str | None = None) -> ParsedDocument:
    """Parse a workbook into one labelled block per row."""
    try:
        import openpyxl
    except ImportError:  # pragma: no cover - dependency is declared
        return ParsedDocument(warnings=("openpyxl is not installed",))

    import io

    try:
        # data_only: read computed values rather than formula source.
        workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        log.warning("could not read workbook: %s", exc)
        return ParsedDocument(warnings=(f"could not read this spreadsheet: {exc}",))

    blocks: list[Block] = []
    warnings: list[str] = []

    try:
        for sheet in workbook.worksheets:
            name = sheet.title
            blocks.append(Block(kind=BlockKind.HEADING, text=name, level=1))
            headers: list[str] = []
            emitted = 0

            for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                if emitted >= MAX_ROWS_PER_SHEET:
                    warnings.append(f"sheet {name!r}: stopped after {MAX_ROWS_PER_SHEET:,} rows")
                    break
                cells = [_render(v) for v in row]
                if not any(cells):
                    continue
                if not headers and looks_like_header_row(cells):
                    headers = cells
                    continue
                text = table_row_text(headers, cells)
                if not text:
                    continue
                blocks.append(
                    Block(
                        kind=BlockKind.TABLE_ROW,
                        text=text,
                        heading_path=(name,),
                        locator=f"{name}!A{index}",
                    )
                )
                emitted += 1
    except Exception as exc:  # noqa: BLE001
        warnings.append(f"stopped reading partway through: {exc}")
    finally:
        with contextlib.suppress(Exception):  # pragma: no cover
            workbook.close()

    return ParsedDocument(blocks=tuple(blocks), title=title, warnings=tuple(warnings))
